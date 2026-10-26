# parse MČR 2026 xlsx exports -> data/*.json
import openpyxl, json, re, sys
D='01_RAW/docs/'
def nm(a,b): return (str(b or '').strip()+' '+str(a or '').strip()).strip()
def day(f):
    ws=openpyxl.load_workbook(D+f,data_only=True).active
    rows=list(ws.iter_rows(values_only=True))
    hdr=rows[4]; title=[c for c in rows[1] if isinstance(c,str) and 'Název' in c][0]
    out=[]
    for r in rows[5:]:
        if not r or r[2] is None: continue
        if not isinstance(r[4+36],int): continue  # prázdné místo (Mirek Dušín = volný los, 39/40)
        rounds=[{'los':r[4+3*i],'fish':r[5+3*i],'pts':r[6+3*i]} for i in range(12)]
        i0=4+36
        out.append({'card':r[0],'name':nm(r[1],r[2]),'country':r[3],'rounds':rounds,'place':r[i0],'points':r[i0+1],'wins':r[i0+2],'draws':r[i0+3],'fish':r[i0+4]})
    return title,out
def rank(f):
    ws=openpyxl.load_workbook(D+f,data_only=True).active
    rows=list(ws.iter_rows(values_only=True))
    events=[c for c in rows[1] if c and c!='HARD' and not str(c).startswith('Filters')]; dates=[c for c in rows[0] if c]
    hdr=rows[2]; n=len(events)
    out=[]
    for r in rows[3:]:
        if not r or r[1] is None: continue
        res=[{'pos':r[3+2*i],'counted':r[4+2*i]} for i in range(n)]
        tail=r[3+2*n:]
        out.append({'n':r[0],'name':nm(r[1],r[2]),'res':res,'visits':tail[0],'sum':tail[1],'avg':tail[2],'rank':tail[3]})
    return {'filter':rows[1][1],'events':events,'dates':dates,'rows':out}
t1,d1=day('tr_b.xlsx'); t2,d2=day('tr_a.xlsx')
assert 'Den 1' in t1 and 'Den 2' in t2
json.dump({'d1':d1,'d2':d2,'mcr':rank('rk_a.xlsx'),'season_all':rank('rk_b.xlsx'),'season_best6':rank('rk_c.xlsx')},open('data/results.json','w'),ensure_ascii=False,indent=1)
print(len(d1),len(d2))
